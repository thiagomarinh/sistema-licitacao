import io


def gerar_excel(data, empresa_info=None):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image as XLImage
    except ImportError:
        raise ImportError('Instale openpyxl: pip install openpyxl')

    config = data.get('config', {'taxa': 7, 'lucro_liq': 9, 'dasn': 3})
    itens = data.get('itens', [])
    titulo = data.get('titulo', 'COMPOSIÇÃO DE PREÇOS')

    taxa = config['taxa'] / 100
    lucro_liq = config['lucro_liq'] / 100
    dasn = config['dasn'] / 100
    divisor = 1 - taxa - lucro_liq - dasn

    empresa_info = empresa_info or {}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Composição de Preços"

    # Estilos
    header_fill = PatternFill("solid", fgColor="1F4E79")
    total_fill = PatternFill("solid", fgColor="D6E4F7")
    timbrado_fill = PatternFill("solid", fgColor="F0F4FA")
    header_font = Font(bold=True, color="FFFFFF", size=9)
    bold_font = Font(bold=True, size=9)
    normal_font = Font(size=9)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center')
    thin = Side(style='thin', color='AAAAAA')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    azul_linha = Side(style='medium', color='1F4E79')

    # ------------------------------------------------------------------ #
    # CABEÇALHO TIMBRADO — ocupa as linhas 1 e 2                          #
    # Estas linhas serão repetidas em cada página impressa                 #
    # ------------------------------------------------------------------ #
    TOTAL_COLS = 18

    # Fundo timbrado linha 1-2
    for r in [1, 2]:
        for c in range(1, TOTAL_COLS + 1):
            ws.cell(row=r, column=c).fill = timbrado_fill

    # Linha azul separadora abaixo do timbrado (linha 3 superior)
    for c in range(1, TOTAL_COLS + 1):
        ws.cell(row=2, column=c).border = Border(bottom=azul_linha)

    ws.row_dimensions[1].height = 40
    ws.row_dimensions[2].height = 20

    # Coluna B–R: nome e dados da empresa (centro)
    ws.merge_cells('B1:R1')
    ws['B1'].value = empresa_info.get('nome', '').upper()
    ws['B1'].font = Font(bold=True, size=14, color="1F4E79")
    ws['B1'].alignment = Alignment(horizontal='center', vertical='center')

    linha2_partes = []
    if empresa_info.get('cnpj'):
        linha2_partes.append(f"CNPJ: {empresa_info['cnpj']}")
    if empresa_info.get('endereco'):
        linha2_partes.append(empresa_info['endereco'])
    if empresa_info.get('telefone'):
        linha2_partes.append(f"Tel: {empresa_info['telefone']}")
    if empresa_info.get('site'):
        linha2_partes.append(empresa_info['site'])

    ws.merge_cells('B2:R2')
    ws['B2'].value = '   |   '.join(linha2_partes)
    ws['B2'].font = Font(size=8, color="444444")
    ws['B2'].alignment = Alignment(horizontal='center', vertical='center')

    # Logo na coluna A (linhas 1-2) se existir
    if empresa_info.get('logo_path'):
        try:
            img = XLImage(empresa_info['logo_path'])
            img.width = 60
            img.height = 55
            ws.add_image(img, 'A1')
        except Exception:
            pass  # logo inválido — continua sem imagem

    ws.column_dimensions['A'].width = 10

    # ------------------------------------------------------------------ #
    # TÍTULO DA COMPOSIÇÃO — linha 4                                       #
    # ------------------------------------------------------------------ #
    ws.merge_cells('A4:S4')
    ws['A4'] = titulo
    ws['A4'].font = Font(bold=True, size=14, color="1F4E79")
    ws['A4'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[4].height = 30

    # Config linha — linha 5
    ws.merge_cells('A5:S5')
    ws['A5'] = (
        f"Taxa: {config['taxa']}%   |   Lucro Líquido: {config['lucro_liq']}%"
        f"   |   DASN: {config['dasn']}%   |   Desconto Total: {config['taxa'] + config['lucro_liq'] + config['dasn']}%"
    )
    ws['A5'].font = Font(italic=True, size=9, color="555555")
    ws['A5'].alignment = Alignment(horizontal='center', vertical='center')

    # Linha em branco 3 (separação visual entre timbrado e título)
    ws.row_dimensions[3].height = 6

    # ------------------------------------------------------------------ #
    # CABEÇALHOS DA TABELA — linha 6                                       #
    # ------------------------------------------------------------------ #
    HEADER_ROW = 6
    headers = [
        '#', 'DESCRIÇÃO', 'UND', 'QUANT', 'FORNECEDOR', 'MARCA SELECIONADA',
        'CUSTO UNIT.\nAQUISIÇÃO', 'CUSTO TOTAL\nAQUISIÇÃO',
        f'PREÇO VENDA\nUNIT. ({100 - round((1 - divisor) * 100)}%)',
        'PREÇO TOTAL\nVENDA', 'LUCRO\nBRUTO %',
        'NOVO\nLANCE', 'NOVO\nTOTAL', 'NOVO\nLUCRO %',
        'NOVO\nLUCRO R$', f'LUCRO LÍQ.\nR$ ({config["lucro_liq"]}%)',
        f'DASN\nR$ ({config["dasn"]}%)', 'CUSTO\nOPERACIONAL'
    ]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    ws.row_dimensions[HEADER_ROW].height = 40

    # ------------------------------------------------------------------ #
    # DADOS                                                                #
    # ------------------------------------------------------------------ #
    totais = {
        'custo_total': 0, 'preco_total': 0, 'novo_total': 0,
        'novo_lucro_rs': 0, 'lucro_liq_rs': 0, 'dasn_rs': 0, 'custo_op': 0
    }

    DATA_START = HEADER_ROW + 1

    for i, item in enumerate(itens):
        row = DATA_START + i
        qtd = float(item.get('quantidade', 0) or 0)
        custo_unit = float(item.get('custo_unit', 0) or 0)
        novo_lance = float(item.get('novo_lance', 0) or 0)

        custo_total = qtd * custo_unit
        preco_venda_unit = (custo_unit / divisor) if divisor > 0 else 0
        preco_total = qtd * preco_venda_unit
        lucro_bruto_pct = ((preco_venda_unit - custo_unit) / custo_unit * 100) if custo_unit > 0 else 0
        novo_total = qtd * novo_lance
        novo_lucro_pct = (((novo_lance - custo_unit) / novo_lance) - dasn - lucro_liq) * 100 if novo_lance > 0 else 0
        novo_lucro_rs = novo_total * (novo_lucro_pct / 100)
        lucro_liq_rs = novo_total * lucro_liq
        dasn_rs = novo_total * dasn
        custo_op = novo_total * taxa

        totais['custo_total'] += custo_total
        totais['preco_total'] += preco_total
        totais['novo_total'] += novo_total
        totais['novo_lucro_rs'] += novo_lucro_rs
        totais['lucro_liq_rs'] += lucro_liq_rs
        totais['dasn_rs'] += dasn_rs
        totais['custo_op'] += custo_op

        values = [
            item.get('num_item', i + 1), item.get('descricao', ''), item.get('unidade', ''), qtd,
            item.get('fornecedor', ''), item.get('marca', ''),
            custo_unit, custo_total, preco_venda_unit, preco_total,
            lucro_bruto_pct / 100, novo_lance, novo_total,
            novo_lucro_pct / 100, novo_lucro_rs, lucro_liq_rs, dasn_rs, custo_op
        ]

        fill = PatternFill("solid", fgColor="EBF3FB") if i % 2 == 1 else None

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = normal_font
            cell.border = border
            if fill:
                cell.fill = fill
            if col in [1, 3]:
                cell.alignment = center
            elif col in [7, 8, 9, 10, 12, 13, 15, 16, 17, 18]:
                cell.number_format = 'R$ #,##0.00'
                cell.alignment = right
            elif col in [11, 14]:
                cell.number_format = '0.00%'
                cell.alignment = center
            elif col == 4:
                cell.number_format = '#,##0.00'
                cell.alignment = center

    # ------------------------------------------------------------------ #
    # TOTAIS                                                               #
    # ------------------------------------------------------------------ #
    total_row = DATA_START + len(itens)
    ws.cell(row=total_row, column=1, value='TOTAIS').font = bold_font
    ws.merge_cells(f'A{total_row}:F{total_row}')
    ws.cell(row=total_row, column=1).alignment = center

    novo_lucro_pct_total = (totais['novo_lucro_rs'] / totais['novo_total'] * 100) if totais['novo_total'] > 0 else 0

    total_values = {
        8: totais['custo_total'], 10: totais['preco_total'],
        13: totais['novo_total'], 14: novo_lucro_pct_total / 100,
        15: totais['novo_lucro_rs'], 16: totais['lucro_liq_rs'],
        17: totais['dasn_rs'], 18: totais['custo_op']
    }

    for col in range(1, TOTAL_COLS + 1):
        cell = ws.cell(row=total_row, column=col)
        if col in total_values:
            cell.value = total_values[col]
        cell.fill = total_fill
        cell.font = bold_font
        cell.border = border
        if col in [8, 10, 13, 15, 16, 17, 18]:
            cell.number_format = 'R$ #,##0.00'
            cell.alignment = right
        elif col == 14:
            cell.number_format = '0.00%'
            cell.alignment = center

    # ------------------------------------------------------------------ #
    # LARGURAS DAS COLUNAS                                                 #
    # ------------------------------------------------------------------ #
    col_widths = [10, 35, 6, 7, 18, 20, 12, 12, 12, 12, 10, 12, 12, 10, 12, 12, 10, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------ #
    # CONFIGURAÇÕES DE IMPRESSÃO                                           #
    # Linhas 1-2 (timbrado) repetem em TODAS as páginas impressas          #
    # ------------------------------------------------------------------ #
    ws.freeze_panes = f'A{DATA_START}'

    ws.print_title_rows = '1:2'   # cabeçalho timbrado repete em cada página

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.print_options.horizontalCentered = True

    # Margens (em polegadas)
    ws.page_margins.top    = 0.5
    ws.page_margins.bottom = 0.7
    ws.page_margins.left   = 0.5
    ws.page_margins.right  = 0.5
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.3

    # ------------------------------------------------------------------ #
    # RODAPÉ — exibido em todas as páginas impressas                       #
    # ------------------------------------------------------------------ #
    nome_empresa = empresa_info.get('nome', '')
    cnpj_empresa = empresa_info.get('cnpj', '')
    rodape_esq = nome_empresa
    if cnpj_empresa:
        rodape_esq += f' | CNPJ: {cnpj_empresa}'

    # Códigos do Excel: &L=esquerda, &C=centro, &R=direita, &P=nro página, &N=total páginas
    ws.oddFooter.left.text  = rodape_esq
    ws.oddFooter.left.size  = 7
    ws.oddFooter.left.color = "444444"
    ws.oddFooter.center.text  = "COMPOSIÇÃO DE CUSTOS - USO INTERNO"
    ws.oddFooter.center.size  = 7
    ws.oddFooter.center.color = "888888"
    ws.oddFooter.right.text  = "Página &P de &N"
    ws.oddFooter.right.size  = 7
    ws.oddFooter.right.color = "444444"

    # Mesmo rodapé para páginas pares
    ws.evenFooter.left.text   = ws.oddFooter.left.text
    ws.evenFooter.left.size   = 7
    ws.evenFooter.left.color  = "444444"
    ws.evenFooter.center.text = ws.oddFooter.center.text
    ws.evenFooter.center.size = 7
    ws.evenFooter.center.color = "888888"
    ws.evenFooter.right.text  = ws.oddFooter.right.text
    ws.evenFooter.right.size  = 7
    ws.evenFooter.right.color = "444444"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
